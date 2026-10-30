import os
import sys
import subprocess
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_notebook():
    repo_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
    docx_path = os.path.join(repo_dir, "notebook.docx")
    pdf_path = os.path.join(repo_dir, "notebook.pdf")

    doc = docx.Document()

    # 1. 页面设定: A4, 边距 1.2cm (~0.47 in)
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.47)
    section.bottom_margin = Inches(0.47)
    section.left_margin = Inches(0.47)
    section.right_margin = Inches(0.47)

    # 页眉与页脚
    footer = section.footer
    f_p = footer.paragraphs[0]
    f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    f_run = f_p.add_run("XCPC Team Notebook · mirai_39")
    f_run.font.name = "Segoe UI"
    f_run.font.size = Pt(8)
    f_run.font.color.rgb = RGBColor(128, 128, 128)

    # 2. 标题区
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    t_run = title_p.add_run("XCPC Team Notebook · 个人板子手册")
    t_run.font.name = "Segoe UI"
    t_run.font.size = Pt(16)
    t_run.font.bold = True
    t_run.font.color.rgb = RGBColor(36, 41, 47)

    meta_p = doc.add_paragraph()
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_p.paragraph_format.space_after = Pt(6)
    m_run = meta_p.add_run("维护者: mirai_39 / mirai0914  |  CCPC / ICPC 冲银  |  纯血 vector · struct 自闭环封装 · 竞赛风")
    m_run.font.name = "Segoe UI"
    m_run.font.size = Pt(8.5)
    m_run.font.color.rgb = RGBColor(87, 96, 106)

    # 编译提示框
    notice_tbl = doc.add_table(rows=1, cols=1)
    notice_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = notice_tbl.cell(0, 0)
    set_cell_background(cell, "F6F8FA")
    set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
    np = cell.paragraphs[0]
    np.paragraph_format.space_before = Pt(0)
    np.paragraph_format.space_after = Pt(0)
    nr1 = np.add_run("【本地防爆栈编译】: ")
    nr1.font.bold = True
    nr1.font.size = Pt(8)
    nr1.font.color.rgb = RGBColor(207, 34, 46)
    nr2 = np.add_run("g++ -O2 -std=c++20 solve.cpp -Wl,--stack=268435456 -o solve.exe\n")
    nr2.font.name = "Consolas"
    nr2.font.size = Pt(8)
    nr2.font.color.rgb = RGBColor(36, 41, 47)
    nr3 = np.add_run("【快速 I/O】: ")
    nr3.font.bold = True
    nr3.font.size = Pt(8)
    nr4 = np.add_run("cin.tie(nullptr)->sync_with_stdio(false);")
    nr4.font.name = "Consolas"
    nr4.font.size = Pt(8)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 板子列表
    boards = [
        {
            "id": "1",
            "title": "欧拉线性筛 (基础素数筛 + 扩展积性函数 phi & mu)",
            "formula": "• φ(p) = p-1; i%p==0 ⇒ φ(ip) = φ(i)*p; 否则 φ(i)*(p-1)\n• μ(1)=1, μ(p)=-1; i%p==0 ⇒ μ(ip) = 0; 否则 -μ(i)",
            "file": os.path.join(repo_dir, "src/math/euler_sieve.cpp")
        },
        {
            "id": "2",
            "title": "快速数论变换 (NTT 多项式乘法 / 快速幂 / 分治连乘)",
            "formula": "• MOD = 998244353, 原根 G = 3\n• mult: O((N+M)log(N+M)); poly_pow: O(L log L log k); poly_prod: O(N log^2 N)",
            "file": os.path.join(repo_dir, "src/math/ntt.cpp")
        },
        {
            "id": "3",
            "title": "分拆数 (五边形数定理 + 根号分治 DP)",
            "formula": "• 广义五边形数: g_k = k(3k±1)/2, 符号循环: + + - - + + - -\n• 根号分治: ≤ B 完全背包; > B 整体+1 与加新数 (B+1) 滚存 DP",
            "file": os.path.join(repo_dir, "src/math/partition.cpp")
        },
        {
            "id": "4",
            "title": "斯特林数 (O(N^2) 递推打表 + NTT 单行求解)",
            "formula": "• 普通幂转下降幂: x^n = ∑ {n, k} x^{underline k}\n• 上升幂转普通幂: x^{overline n} = ∑ [n, k] x^k = ∏_{i=0}^{n-1} (x+i)\n• 卷积容斥: {n, k} = ∑_{j=0}^k (-1)^{k-j}/(k-j)! * (j^n / j!)",
            "file": os.path.join(repo_dir, "src/math/stirling.cpp")
        },
        {
            "id": "5",
            "title": "珂朵莉树 (ODT - Old Driver Tree / Chtholly Tree)",
            "formula": "• 核心避坑准则: 双端点切分必须【先 split(r + 1)，再 split(l)】！否则左端点迭代器可能失效导致 RE！\n• 结构 set<array<int, 3>>，在频繁随机区间推平 assign 下块数期望 O(log log n)",
            "file": os.path.join(repo_dir, "src/data_structure/odt.cpp")
        },
        {
            "id": "6",
            "title": "主席树 (可持久化权值线段树 / 静态区间第 k 小 / 树上第 k 小)",
            "formula": "• 基础: roots[r] 与 roots[l-1] 差分二分 | 单次插/查 O(log M), 空间 O(N log M) 动态开点\n• 扩展: 区间 ≤ x 计数 count_le; 树上路径差分 u + v - lca - fa_lca",
            "file": os.path.join(repo_dir, "src/data_structure/persistent_seg_tree.cpp")
        }
    ]

    for b in boards:
        h_p = doc.add_paragraph()
        h_p.paragraph_format.space_before = Pt(8)
        h_p.paragraph_format.space_after = Pt(2)
        h_p.paragraph_format.keep_with_next = True
        hr = h_p.add_run(f"{b['id']}. {b['title']}")
        hr.font.name = "Segoe UI"
        hr.font.size = Pt(11)
        hr.font.bold = True
        hr.font.color.rgb = RGBColor(9, 105, 218)

        if b.get("formula"):
            f_tbl = doc.add_table(rows=1, cols=1)
            f_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            f_cell = f_tbl.cell(0, 0)
            set_cell_background(f_cell, "FFF8C5")
            set_cell_margins(f_cell, top=60, bottom=60, left=100, right=100)
            fp = f_cell.paragraphs[0]
            fp.paragraph_format.space_before = Pt(0)
            fp.paragraph_format.space_after = Pt(0)
            fr = fp.add_run(b["formula"])
            fr.font.name = "Consolas"
            fr.font.size = Pt(7.8)
            fr.font.color.rgb = RGBColor(115, 92, 15)

        code = ""
        if os.path.exists(b["file"]):
            with open(b["file"], "r", encoding="utf-8") as f:
                code = f.read().strip()

        c_tbl = doc.add_table(rows=1, cols=1)
        c_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        c_cell = c_tbl.cell(0, 0)
        set_cell_background(c_cell, "F6F8FA")
        set_cell_margins(c_cell, top=80, bottom=80, left=100, right=100)
        cp = c_cell.paragraphs[0]
        cp.paragraph_format.space_before = Pt(0)
        cp.paragraph_format.space_after = Pt(0)
        cp.paragraph_format.line_spacing = 1.15
        cr = cp.add_run(code)
        cr.font.name = "Consolas"
        cr.font.size = Pt(7.4)
        cr.font.color.rgb = RGBColor(36, 41, 47)

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    doc.save(docx_path)
    print(f"[OK] Generated Word document: {docx_path}")

    # 调用 Word 导出真正标准的 PDF
    ps_cmd = f"""
    $docxPath = "{docx_path}"
    $pdfPath = "{pdf_path}"
    try {{
        $word = New-Object -ComObject Word.Application
        $word.Visible = $false
        $doc = $word.Documents.Open($docxPath)
        $doc.SaveAs([ref]$pdfPath, [ref]17)
        $doc.Close()
        $word.Quit()
        Write-Output "[OK] Generated PDF: $pdfPath"
    }} catch {{
        Write-Error $_.Exception.Message
    }}
    """
    res = subprocess.run(["powershell", "-ExecutionPolicy", "Bypass", "-Command", ps_cmd], capture_output=True, text=True)
    if res.stdout:
        print(res.stdout.strip())
    if res.stderr:
        print(res.stderr.strip())

if __name__ == "__main__":
    create_notebook()
